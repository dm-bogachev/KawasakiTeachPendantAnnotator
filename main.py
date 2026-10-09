import sys, json, os
from dataclasses import dataclass, asdict
from pathlib import Path

from PySide6.QtCore import Qt, QRect, QRectF, Signal
from PySide6.QtGui import QPixmap, QPainter, QPen, QColor, QImage, QGuiApplication
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QLabel, QPushButton, QFileDialog,
    QMessageBox, QVBoxLayout, QHBoxLayout, QFormLayout, QLineEdit, QTextEdit,
    QDialog, QDialogButtonBox, QCheckBox, QSpinBox, QGroupBox, QSplitter,
    QScrollArea, QStatusBar, QToolButton, QColorDialog, QSlider
)
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


GRID = {
    # Individual cell hitboxes measured from the supplied 642x512 screenshot.
    "x_cells": [(11, 90), (101, 181), (191, 271), (281, 361), (371, 451), (461, 541), (551, 631)],
    "y_rows": [(191, 271), (271, 351), (351, 431), (431, 510)],
    "cols": 7, "rows": 4
}
APP_VERSION = 1


@dataclass
class PanelElement:
    key: str
    kind: str
    rect: list
    number: str = ""
    english_name: str = ""
    description: str = ""


class ElementDialog(QDialog):
    def __init__(self, element, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Параметры элемента")
        self.resize(390, 250)
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.number = QLineEdit(element.number)
        self.name = QLineEdit(element.english_name)
        self.name.setPlaceholderText("Например: Program window")
        self.description = QTextEdit(element.description)
        self.description.setPlaceholderText("Комментарий / описание на русском")
        self.description.setFixedHeight(90)
        form.addRow("Номер:", self.number)
        form.addRow("Имя на английском:", self.name)
        form.addRow("Описание:", self.description)
        layout.addLayout(form)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def values(self):
        return self.number.text().strip(), self.name.text().strip(), self.description.toPlainText().strip()


class ImageCanvas(QWidget):
    elementClicked = Signal(str)
    customRectCreated = Signal(list)
    selectionChanged = Signal(str)

    def __init__(self):
        super().__init__()
        self.setMinimumSize(642, 512)
        self.pixmap = QPixmap()
        self.elements = {}
        self.selected_key = None
        self.draw_numbers = True
        self.custom_mode = False
        self.drag_start = None
        self.drag_end = None
        self.setMouseTracking(True)

    def set_image(self, path):
        pix = QPixmap(path)
        if pix.isNull():
            return False
        self.pixmap = pix
        self.updateGeometry()
        self.update()
        return True

    def set_pixmap(self, pix):
        if pix.isNull():
            return False
        self.pixmap = pix
        self.updateGeometry()
        self.update()
        return True

    def sizeHint(self):
        if not self.pixmap.isNull():
            return self.pixmap.size()
        return super().sizeHint()

    def image_rect(self):
        if self.pixmap.isNull():
            return QRectF(0, 0, 642, 512)
        # Keep native image size for accurate 642x512 coordinate mapping.
        return QRectF(0, 0, self.pixmap.width(), self.pixmap.height())

    def coordinate(self, pos):
        r = self.image_rect()
        x, y = pos.x(), pos.y()
        if x < r.left() or y < r.top() or x >= r.right() or y >= r.bottom():
            return None
        return int(x), int(y)

    def grid_key_at(self, x, y):
        g = GRID
        col = next((i for i, (x1, x2) in enumerate(g["x_cells"]) if x1 <= x < x2), None)
        row = next((i for i, (y1, y2) in enumerate(g["y_rows"]) if y1 <= y < y2), None)
        if col is None or row is None:
            return None
        return f"cell_{row+1}_{col+1}"

    def mousePressEvent(self, event):
        if event.button() != Qt.LeftButton or self.pixmap.isNull():
            return
        p = self.coordinate(event.position().toPoint())
        if p is None:
            return
        x, y = p
        if self.custom_mode:
            self.drag_start = (x, y)
            self.drag_end = (x, y)
            self.update()
            return
        key = self.grid_key_at(x, y)
        if key:
            self.elementClicked.emit(key)
            return
        # Select an existing custom/top element by clicking inside it.
        for k, e in reversed(list(self.elements.items())):
            if e.kind == "custom":
                rx, ry, rw, rh = e.rect
                if rx <= x <= rx + rw and ry <= y <= ry + rh:
                    self.elementClicked.emit(k)
                    return
        self.selectionChanged.emit("Кликните ячейку сетки или используйте кнопку «Кастомный элемент».")

    def mouseMoveEvent(self, event):
        if self.drag_start is not None:
            p = self.coordinate(event.position().toPoint())
            if p:
                self.drag_end = p
                self.update()

    def mouseReleaseEvent(self, event):
        if self.drag_start is None:
            return
        p = self.coordinate(event.position().toPoint())
        if p:
            self.drag_end = p
        x1, y1 = self.drag_start
        x2, y2 = self.drag_end or self.drag_start
        self.drag_start = self.drag_end = None
        self.custom_mode = False
        self.customRectCreated.emit([min(x1,x2), min(y1,y2), abs(x2-x1), abs(y2-y1)])
        self.update()

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.fillRect(self.rect(), QColor("#eeeeee"))
        if self.pixmap.isNull():
            painter.drawText(self.rect(), Qt.AlignCenter, "Загрузите изображение панели (рекомендуется 642×512)")
            return
        painter.drawPixmap(0, 0, self.pixmap)
        painter.setRenderHint(QPainter.Antialiasing, False)
        # Show grid hit areas only as subtle overlay.
        g = GRID
        painter.setPen(QPen(QColor(30, 120, 230, 170), 1, Qt.DashLine))
        for x1, x2 in g["x_cells"]:
            for y1, y2 in g["y_rows"]:
                painter.drawRect(QRect(x1, y1, x2-x1, y2-y1))
        for key, el in self.elements.items():
            x,y,w,h = el.rect
            rect = QRect(x,y,w,h)
            if key == self.selected_key:
                color = QColor(255, 170, 0, 75)
                pen = QPen(QColor("#ff8c00"), 2)
            else:
                color = QColor(0, 160, 90, 55)
                pen = QPen(QColor("#00965a"), 1)
            painter.fillRect(rect, color)
            painter.setPen(pen)
            painter.drawRect(rect)
            label = el.number or ("?" if el.kind == "custom" else "")
            if self.draw_numbers and label:
                font = painter.font()
                font.setPointSize(getattr(self, "number_font_size", 12))
                painter.setFont(font)
                metrics = painter.fontMetrics()
                label_w = metrics.horizontalAdvance(label) + 10
                label_h = metrics.height() + 4
                bg = QColor(getattr(self, "number_bg_color", "#ffffff"))
                bg.setAlpha(getattr(self, "number_bg_opacity", 225))
                painter.fillRect(QRect(x+2, y+2, label_w, label_h), bg)
                painter.setPen(QPen(QColor(getattr(self, "number_text_color", "#111111"))))
                painter.drawText(x+6, y+2+metrics.ascent()+2, label)
        if self.drag_start and self.drag_end:
            x1,y1 = self.drag_start
            x2,y2 = self.drag_end
            painter.setPen(QPen(QColor("#d02020"), 2, Qt.DashLine))
            painter.drawRect(QRect(min(x1,x2),min(y1,y2),abs(x2-x1),abs(y2-y1)))


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Robot Panel Annotator")
        self.resize(1120, 760)
        self.image_path = None
        self.elements = {}
        self.custom_counter = 0
        self.include_cell_images = True
        # Styles for number labels, persisted in the project file.
        self.number_font_size = 12
        self.number_text_color = "#111111"
        self.number_bg_color = "#ffffff"
        self.number_bg_opacity = 225  # 0..255
        self.canvas = ImageCanvas()
        self.canvas.number_font_size = self.number_font_size
        self.canvas.number_text_color = self.number_text_color
        self.canvas.number_bg_color = self.number_bg_color
        self.canvas.number_bg_opacity = self.number_bg_opacity
        self.canvas.elementClicked.connect(self.edit_element)
        self.canvas.customRectCreated.connect(self.add_custom_element)
        self.canvas.selectionChanged.connect(self.statusBar().showMessage)

        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        toolbar = QHBoxLayout()
        self.btn_new = QPushButton("Новый проект")
        self.btn_open = QPushButton("Загрузить изображение…")
        self.btn_clip = QPushButton("Вставить из буфера")
        self.btn_save = QPushButton("Сохранить проект…")
        self.btn_load = QPushButton("Открыть проект…")
        self.btn_add = QPushButton("Кастомный элемент")
        self.btn_export = QPushButton("Экспорт в Word…")
        for b in (self.btn_new, self.btn_open, self.btn_clip, self.btn_save, self.btn_load, self.btn_add, self.btn_export):
            toolbar.addWidget(b)
        root.addLayout(toolbar)

        splitter = QSplitter()
        scroll = QScrollArea()
        scroll.setWidget(self.canvas)
        scroll.setWidgetResizable(False)
        splitter.addWidget(scroll)
        side = QWidget()
        side_layout = QVBoxLayout(side)
        info = QLabel(
            "Как пользоваться:\n"
            "• Кликните по ячейке сетки 7×4, чтобы задать параметры.\n"
            "• Для дополнительных элементов нажмите «Кастомный элемент» и обведите область мышью.\n"
            "• Поля можно оставить пустыми.\n"
            "• Изображение встраивается в проект JSON в Base64.\n\n"
            "Координаты сетки заданы для исходного скриншота 642×512."
        )
        info.setWordWrap(True)
        side_layout.addWidget(info)
        self.chk_images = QCheckBox("Добавлять изображения ячеек в таблицу Word")
        self.chk_images.setChecked(True)
        side_layout.addWidget(self.chk_images)
        self.chk_numbers = QCheckBox("Показывать номера на изображении")
        self.chk_numbers.setChecked(True)
        self.chk_numbers.toggled.connect(self.toggle_numbers)
        side_layout.addWidget(self.chk_numbers)
        style_group = QGroupBox("Оформление нумерации")
        style_layout = QFormLayout(style_group)
        self.spin_font_size = QSpinBox()
        self.spin_font_size.setRange(6, 48)
        self.spin_font_size.setValue(self.number_font_size)
        style_layout.addRow("Размер шрифта:", self.spin_font_size)

        self.btn_text_color = QPushButton("Выбрать цвет текста…")
        self.lbl_text_color = QLabel(self.number_text_color)
        text_row = QHBoxLayout()
        text_row.addWidget(self.btn_text_color)
        text_row.addWidget(self.lbl_text_color)
        style_layout.addRow("Цвет текста:", text_row)

        self.btn_bg_color = QPushButton("Выбрать цвет фона…")
        self.lbl_bg_color = QLabel(self.number_bg_color)
        bg_row = QHBoxLayout()
        bg_row.addWidget(self.btn_bg_color)
        bg_row.addWidget(self.lbl_bg_color)
        style_layout.addRow("Цвет фона:", bg_row)

        self.slider_bg_opacity = QSlider(Qt.Horizontal)
        self.slider_bg_opacity.setRange(0, 255)
        self.slider_bg_opacity.setValue(self.number_bg_opacity)
        self.lbl_bg_opacity = QLabel(f"{round(self.number_bg_opacity / 255 * 100)}%")
        opacity_row = QHBoxLayout()
        opacity_row.addWidget(self.slider_bg_opacity)
        opacity_row.addWidget(self.lbl_bg_opacity)
        style_layout.addRow("Прозрачность фона:", opacity_row)
        side_layout.addWidget(style_group)

        self.lbl_selected = QLabel("Выбранный элемент: —")
        self.lbl_selected.setWordWrap(True)
        side_layout.addWidget(self.lbl_selected)
        self.btn_edit = QPushButton("Изменить выбранный элемент")
        self.btn_delete = QPushButton("Удалить выбранный элемент")
        side_layout.addWidget(self.btn_edit)
        side_layout.addWidget(self.btn_delete)
        side_layout.addStretch(1)
        splitter.addWidget(side)
        splitter.setSizes([760, 330])
        root.addWidget(splitter)

        self.btn_new.clicked.connect(self.new_project)
        self.btn_open.clicked.connect(self.open_image)
        self.btn_clip.clicked.connect(self.paste_image)
        self.btn_save.clicked.connect(self.save_project)
        self.btn_load.clicked.connect(self.load_project)
        self.btn_add.clicked.connect(self.start_custom)
        self.btn_export.clicked.connect(self.export_word)
        self.btn_edit.clicked.connect(self.edit_selected)
        self.btn_delete.clicked.connect(self.delete_selected)
        self.chk_images.toggled.connect(lambda v: setattr(self, "include_cell_images", v))
        self.spin_font_size.valueChanged.connect(self.set_number_font_size)
        self.btn_text_color.clicked.connect(self.choose_text_color)
        self.btn_bg_color.clicked.connect(self.choose_bg_color)
        self.slider_bg_opacity.valueChanged.connect(self.set_bg_opacity)
        self.statusBar().showMessage("Готово. Загрузите изображение панели.")

    def new_project(self):
        answer = QMessageBox.question(
            self, "Новый проект",
            "Создать новый проект? Несохранённые изменения будут потеряны.",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if answer != QMessageBox.Yes:
            return
        self.image_path = None
        self.elements = {}
        self.custom_counter = 0
        self.canvas.pixmap = QPixmap()
        self.canvas.elements = self.elements
        self.canvas.selected_key = None
        self.canvas.custom_mode = False
        self.canvas.drag_start = self.canvas.drag_end = None
        self.lbl_selected.setText("Выбранный элемент: —")
        self.chk_images.setChecked(True)
        self.chk_numbers.setChecked(True)
        self.canvas.updateGeometry()
        self.canvas.update()
        self.statusBar().showMessage("Создан новый проект.")

    def set_number_font_size(self, value):
        self.number_font_size = value
        self.canvas.number_font_size = value
        self.canvas.update()

    def choose_text_color(self):
        color = QColorDialog.getColor(QColor(self.number_text_color), self, "Цвет текста номера")
        if color.isValid():
            self.number_text_color = color.name()
            self.lbl_text_color.setText(color.name())
            self.canvas.number_text_color = color.name()
            self.canvas.update()

    def choose_bg_color(self):
        color = QColorDialog.getColor(QColor(self.number_bg_color), self, "Цвет фона номера")
        if color.isValid():
            self.number_bg_color = color.name()
            self.lbl_bg_color.setText(color.name())
            self.canvas.number_bg_color = color.name()
            self.canvas.update()

    def set_bg_opacity(self, value):
        self.number_bg_opacity = value
        self.canvas.number_bg_opacity = value
        self.lbl_bg_opacity.setText(f"{round(value / 255 * 100)}%")
        self.canvas.update()

    def default_rect(self, key):
        if key.startswith("cell_"):
            _, row, col = key.split("_")
            g = GRID
            x, x2 = g["x_cells"][int(col)-1]
            y, y2 = g["y_rows"][int(row)-1]
            return [x, y, x2-x, y2-y]
        return [0,0,0,0]

    def open_image(self):
        path, _ = QFileDialog.getOpenFileName(self, "Выберите изображение", "", "Изображения (*.png *.jpg *.jpeg *.bmp *.webp)")
        if path:
            if self.canvas.set_image(path):
                self.image_path = os.path.abspath(path)
                self.statusBar().showMessage(f"Изображение загружено: {Path(path).name}")
            else:
                QMessageBox.warning(self, "Ошибка", "Не удалось открыть изображение.")

    def paste_image(self):
        pix = QGuiApplication.clipboard().pixmap()
        if pix.isNull():
            image = QGuiApplication.clipboard().image()
            if not image.isNull():
                pix = QPixmap.fromImage(image)
        if pix.isNull():
            QMessageBox.information(self, "Буфер обмена", "В буфере обмена не найдено изображение.")
            return
        self.canvas.set_pixmap(pix)
        self.image_path = None
        self.statusBar().showMessage(f"Изображение вставлено из буфера ({pix.width()}×{pix.height()}); при сохранении проекта оно будет встроено.")

    def edit_element(self, key):
        if key not in self.elements:
            kind = "grid" if key.startswith("cell_") else "custom"
            self.elements[key] = PanelElement(key, kind, self.default_rect(key))
        self.canvas.selected_key = key
        self.canvas.elements = self.elements
        dlg = ElementDialog(self.elements[key], self)
        self.lbl_selected.setText(f"Выбранный элемент: {key}")
        if dlg.exec() == QDialog.Accepted:
            number, name, desc = dlg.values()
            e = self.elements[key]
            e.number, e.english_name, e.description = number, name, desc
        self.canvas.update()

    def start_custom(self):
        if self.canvas.pixmap.isNull():
            QMessageBox.information(self, "Нет изображения", "Сначала загрузите изображение.")
            return
        self.canvas.custom_mode = True
        self.statusBar().showMessage("Зажмите левую кнопку мыши и обведите нужную область прямоугольником.")

    def add_custom_element(self, rect):
        if rect[2] < 3 or rect[3] < 3:
            self.statusBar().showMessage("Область слишком маленькая; попробуйте ещё раз.")
            return
        self.custom_counter += 1
        key = f"custom_{self.custom_counter}"
        self.elements[key] = PanelElement(key, "custom", rect)
        self.canvas.elements = self.elements
        self.canvas.selected_key = key
        self.edit_element(key)

    def edit_selected(self):
        if self.canvas.selected_key and self.canvas.selected_key in self.elements:
            self.edit_element(self.canvas.selected_key)
        else:
            QMessageBox.information(self, "Нет выбора", "Сначала выберите элемент на изображении.")

    def delete_selected(self):
        key = self.canvas.selected_key
        if key and key in self.elements:
            del self.elements[key]
            self.canvas.selected_key = None
            self.canvas.elements = self.elements
            self.lbl_selected.setText("Выбранный элемент: —")
            self.canvas.update()

    def toggle_numbers(self, checked):
        self.canvas.draw_numbers = checked
        self.canvas.update()

    def project_data(self):
        image_data = None
        if not self.canvas.pixmap.isNull():
            import base64
            from PySide6.QtCore import QBuffer, QByteArray
            ba = QByteArray()
            buf = QBuffer(ba)
            buf.open(QBuffer.WriteOnly)
            self.canvas.pixmap.save(buf, "PNG")
            image_data = base64.b64encode(bytes(ba)).decode("ascii")
        return {
            "app_version": APP_VERSION,
            "image_path": self.image_path,
            "image_embedded_png_base64": image_data,
            "elements": [asdict(e) for e in self.elements.values()],
            "include_cell_images": self.chk_images.isChecked(),
            "draw_numbers": self.chk_numbers.isChecked(),
            "custom_counter": self.custom_counter,
            "number_font_size": self.number_font_size,
            "number_text_color": self.number_text_color,
            "number_bg_color": self.number_bg_color,
            "number_bg_opacity": self.number_bg_opacity,
        }

    def save_project(self):
        path, _ = QFileDialog.getSaveFileName(self, "Сохранить проект", "robot_panel_project.json", "Проект (*.json)")
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(self.project_data(), f, ensure_ascii=False, indent=2)
            self.statusBar().showMessage(f"Проект сохранён: {path}")
        except Exception as exc:
            QMessageBox.critical(self, "Ошибка сохранения", str(exc))

    def load_project(self):
        path, _ = QFileDialog.getOpenFileName(self, "Открыть проект", "", "Проект (*.json)")
        if not path:
            return
        try:
            with open(path, encoding="utf-8") as f:
                data = json.load(f)
            self.elements = {e["key"]: PanelElement(**e) for e in data.get("elements", [])}
            self.custom_counter = data.get("custom_counter", 0)
            self.number_font_size = int(data.get("number_font_size", 12))
            self.number_text_color = data.get("number_text_color", "#111111")
            self.number_bg_color = data.get("number_bg_color", "#ffffff")
            self.number_bg_opacity = int(data.get("number_bg_opacity", 225))
            self.spin_font_size.setValue(self.number_font_size)
            self.lbl_text_color.setText(self.number_text_color)
            self.lbl_bg_color.setText(self.number_bg_color)
            self.slider_bg_opacity.setValue(self.number_bg_opacity)
            self.lbl_bg_opacity.setText(f"{round(self.number_bg_opacity / 255 * 100)}%")
            self.canvas.number_font_size = self.number_font_size
            self.canvas.number_text_color = self.number_text_color
            self.canvas.number_bg_color = self.number_bg_color
            self.canvas.number_bg_opacity = self.number_bg_opacity
            img_path = data.get("image_path")
            loaded = False
            if img_path and os.path.exists(img_path):
                loaded = self.canvas.set_image(img_path)
                self.image_path = os.path.abspath(img_path) if loaded else None
            if not loaded and data.get("image_embedded_png_base64"):
                import base64
                raw = base64.b64decode(data["image_embedded_png_base64"])
                pix = QPixmap()
                pix.loadFromData(raw, "PNG")
                loaded = self.canvas.set_pixmap(pix)
                self.image_path = None
            if not loaded:
                QMessageBox.warning(self, "Изображение не найдено", "Проект открыт, но исходная картинка недоступна. Загрузите её заново.")
            self.chk_images.setChecked(data.get("include_cell_images", True))
            self.chk_numbers.setChecked(data.get("draw_numbers", True))
            self.canvas.elements = self.elements
            self.canvas.selected_key = None
            self.canvas.update()
            self.statusBar().showMessage(f"Проект загружен: {path}")
        except Exception as exc:
            QMessageBox.critical(self, "Ошибка открытия проекта", str(exc))

    def crop_element(self, el):
        x,y,w,h = el.rect
        rect = QRect(max(0,x), max(0,y), max(1,w), max(1,h))
        return self.canvas.pixmap.copy(rect)

    def export_word(self):
        if self.canvas.pixmap.isNull():
            QMessageBox.information(self, "Нет изображения", "Сначала загрузите изображение панели.")
            return
        path, _ = QFileDialog.getSaveFileName(self, "Экспорт в Word", "robot_panel_reference.docx", "Документ Word (*.docx)")
        if not path:
            return
        try:
            # Render a clean annotated panel with numbers.
            annotated = QPixmap(self.canvas.pixmap.size())
            annotated.fill(Qt.white)
            painter = QPainter(annotated)
            painter.drawPixmap(0,0,self.canvas.pixmap)
            for e in self.elements.values():
                x,y,w,h = e.rect
                painter.setPen(QPen(QColor("#e87500"), 2))
                painter.fillRect(QRect(x,y,w,h), QColor(255,170,0,45))
                painter.drawRect(QRect(x,y,w,h))
                if e.number:
                    font = painter.font()
                    font.setPointSize(self.number_font_size)
                    painter.setFont(font)
                    metrics = painter.fontMetrics()
                    label_w = metrics.horizontalAdvance(e.number) + 10
                    label_h = metrics.height() + 4
                    bg = QColor(self.number_bg_color)
                    bg.setAlpha(self.number_bg_opacity)
                    painter.fillRect(QRect(x+2, y+2, label_w, label_h), bg)
                    painter.setPen(QPen(QColor(self.number_text_color)))
                    painter.drawText(x+6, y+2+metrics.ascent()+2, e.number)
            painter.end()
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                annotated.save(tmp.name, "PNG")
                annotated_path = tmp.name

            doc = Document()
            doc.add_heading("Справочник элементов пульта робота", 0)
            p = doc.add_paragraph("Изображение панели с обозначениями")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_picture(annotated_path, width=Inches(6.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_heading("Таблица элементов", level=1)
            entries = [e for e in self.elements.values() if e.number or e.english_name or e.description]
            entries.sort(key=lambda e: (0, int(e.number)) if e.number.isdigit() else (1, e.number.lower(), e.key))
            headers = (["№", "Изображение", "English name", "Описание / комментарий"] if self.chk_images.isChecked()
                       else ["№", "English name", "Описание / комментарий"])
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i,h in enumerate(headers):
                table.rows[0].cells[i].text = h
            for e in entries:
                cells = table.add_row().cells
                idx = 0
                cells[idx].text = e.number
                idx += 1
                if self.chk_images.isChecked():
                    cells[idx].text = ""
                    cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    crop = self.crop_element(e)
                    if not crop.isNull():
                        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                            crop.save(tmp.name, "PNG")
                            crop_path = tmp.name
                        run = cells[idx].paragraphs[0].add_run()
                        run.add_picture(crop_path, width=Inches(0.9))
                    idx += 1
                cells[idx].text = e.english_name
                cells[idx+1].text = e.description
                for cell in cells:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for section in doc.sections:
                section.top_margin = Inches(0.55)
                section.bottom_margin = Inches(0.55)
            doc.save(path)
            try:
                os.unlink(annotated_path)
            except Exception:
                pass
            QMessageBox.information(self, "Экспорт завершён", f"Документ сохранён:\n{path}")
        except Exception as exc:
            QMessageBox.critical(self, "Ошибка экспорта", f"Не удалось создать Word-документ:\n{exc}")


def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
